"""
US Volume Surge backtest -> Excel report with every buy / sell, grouped by ticker.

Runs the Scenario A strategy on the US universe (tracked tickers + S&P 500 / Russell 1000):
    surge day (volume > N x 10-day avg, close > MA200, MA50 > MA200, up day)
    -> first RED candle after it (drop day)
    -> buy at the surge-day close level within 5 trading days
    -> sell half at +15% (close), sell the rest when a close is below MA21 and the next open is still below.
then replays the signals through a capital account (start capital, N slots, each buy = base stake plus an
even share of the profit made so far) so only the trades that account could actually afford are "taken".

Buy fills are realistic (see ENTRY_MODE in study_volume_surge_2026_04.py):
    stop   (default) buy on the way up: pay max(open, level) on the first day High >= level
    limit  true buy limit: needs Low <= level, pay min(open, level)
    optimistic  the old, over-generous rule (fill at the level even if the stock never traded there)

Sheets: Summary | By Ticker (collapsible groups, one block per ticker with each buy/sell) |
        By Buy Date (the same detail, one block per trade in the order you bought) | Ticker Summary |
        Trades (one row per trade, sorted by buy date) | Transactions (chronological) | Skipped Signals | Still Open

Run from the project root:
    python scripts/us_backtest_report.py
    python scripts/us_backtest_report.py --entry-mode limit --vol-mult 4 --capital 20000 --position 1000 --slots 20
"""
import argparse
import contextlib
import io
import json
import os
import sqlite3
import sys

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import study_volume_surge_2026_04 as study
import capital_simulation as capsim
from study_entry_modes import months

OUT_ROOT = "us_backtest_reports"
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TICKER_FILL = PatternFill("solid", fgColor="DCE6F1")
GREEN = Font(color="15803D")
RED = Font(color="B91C1C")


def generate_trades(by_ticker, mode, vol_mult, first, last, out_dir):
    study.ENTRY_MODE, study.VOL_MULT = mode, vol_mult
    os.makedirs(out_dir, exist_ok=True)
    for f in os.listdir(out_dir):
        if f.startswith("study_volume_surge_") and f.endswith(".json"):
            os.remove(os.path.join(out_dir, f))
    total = 0
    for ym in months(first, last):
        study.SURGE_START, study.SURGE_END = f"{ym}-01", f"{ym}-31"
        trades = []
        for ticker, series in by_ticker.items():
            for surge_idx in study.find_surge_indices(series):
                drop_idx = study.find_drop_day(series, surge_idx)
                if drop_idx is None:
                    continue
                a_idx = study.find_scenario_a(series, surge_idx, drop_idx)
                if a_idx is None:
                    continue
                level = series[surge_idx]["close"]
                trades.append(study.build_trade(
                    ticker, series[surge_idx]["date"], level, series[drop_idx]["date"], "A", series, a_idx,
                    entry_price=study.entry_fill(series, a_idx, level)))
        with open(os.path.join(out_dir, f"study_volume_surge_{ym.replace('-', '_')}.json"), "w") as f:
            json.dump({"trades": trades}, f, indent=2)
        total += len(trades)
    return total


def transaction_rows(t):
    """Buy / sell rows for one executed trade (a row of the 'Trades' dataframe)."""
    rows = [{"date": t.buy_date, "action": "BUY", "price": t.buy_price, "shares": t.shares_bought,
             "amount": -t.invested,
             "note": f"surge {t.surge_date}, drop day {t.drop_date}"}]
    if pd.notna(t.partial_sell_date):
        rows.append({"date": t.partial_sell_date, "action": "SELL HALF (+15% target)", "price": t.partial_sell_price,
                     "shares": t.partial_sell_shares, "amount": t.partial_sell_proceeds, "note": "half sold at the target"})
    closed = t.status == "closed"
    rows.append({"date": t.final_sell_date,
                 "action": ("SELL REST" if pd.notna(t.partial_sell_date) else "SELL ALL") if closed else "STILL OPEN (marked at last close)",
                 "price": t.final_sell_price, "shares": t.final_sell_shares, "amount": t.final_sell_proceeds,
                 "note": (t.exit_type or "") if closed else "not sold yet - value at the last close"})
    return rows


def write_by_ticker(ws, trades):
    headers = ["Ticker / trade", "Date", "Action", "Price", "Shares", "Amount ($)", "Trade P/L ($)", "Trade P/L (%)", "Days held", "Note"]
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.sheet_properties.outlinePr.summaryBelow = False          # group header sits ABOVE its rows
    for ticker, g in trades.groupby("ticker", sort=True):
        pnl = g.total_pnl.sum()
        wins = int((g.total_pnl > 0).sum())
        ws.append([ticker, f"{len(g)} trade(s)", f"{wins} win / {len(g) - wins} loss", None, None, None,
                   round(pnl, 2), round(100 * pnl / g.invested.sum(), 2), None, "ticker total"])
        r = ws.max_row
        for c in ws[r]:
            c.font, c.fill = BOLD, TICKER_FILL
        ws.cell(r, 7).font = Font(bold=True, color="15803D" if pnl >= 0 else "B91C1C")
        for n, t in enumerate(g.sort_values("buy_date").itertuples(), 1):
            days = (pd.to_datetime(t.final_sell_date) - pd.to_datetime(t.buy_date)).days
            for i, row in enumerate(transaction_rows(t)):
                last = i == len(transaction_rows(t)) - 1
                ws.append([f"   trade {n}" if i == 0 else None, row["date"], row["action"], row["price"], row["shares"],
                           round(row["amount"], 2), round(t.total_pnl, 2) if last else None,
                           round(t.total_pnl_pct, 2) if last else None, days if last else None, row["note"]])
                ws.row_dimensions[ws.max_row].outlineLevel = 1
                if last:
                    ws.cell(ws.max_row, 7).font = GREEN if t.total_pnl >= 0 else RED
                    ws.cell(ws.max_row, 8).font = GREEN if t.total_pnl >= 0 else RED
    for col, w in zip("ABCDEFGHIJ", (16, 12, 30, 10, 12, 13, 14, 13, 10, 44)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "0.0000"
        row[4].number_format = "#,##0.0000"
        for i in (5, 6):
            row[i].number_format = "#,##0.00;[Red]-#,##0.00"
        row[7].number_format = "0.00"


def write_by_buy_date(ws, trades):
    """One collapsible block per trade, in buy-date order: a bold header (buy date, ticker, trade result, running
    total of trade P/L) followed by that trade's BUY / SELL rows."""
    headers = ["Buy date / trade", "Ticker", "Action", "Date", "Price", "Shares", "Amount ($)", "Trade P/L ($)",
               "Trade P/L (%)", "Days held", "Running total P/L ($)", "Note"]
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.sheet_properties.outlinePr.summaryBelow = False
    running = 0.0
    for n, t in enumerate(trades.sort_values(["buy_date", "ticker"]).itertuples(), 1):
        running += t.total_pnl
        days = (pd.to_datetime(t.final_sell_date) - pd.to_datetime(t.buy_date)).days
        state = "still open" if t.status != "closed" else ("win" if t.total_pnl > 0 else "loss")
        ws.append([t.buy_date, t.ticker, f"trade #{n} ({state})", None, None, None, None, round(t.total_pnl, 2),
                   round(t.total_pnl_pct, 2), days, round(running, 2), f"invested ${t.invested:,.2f}"])
        r = ws.max_row
        for c in ws[r]:
            c.font, c.fill = BOLD, TICKER_FILL
        col = "15803D" if t.total_pnl >= 0 else "B91C1C"
        for i in (8, 9):
            ws.cell(r, i).font = Font(bold=True, color=col)
        rows = transaction_rows(t)
        for row in rows:
            ws.append([None, None, row["action"], row["date"], row["price"], row["shares"], round(row["amount"], 2),
                       None, None, None, None, row["note"]])
            ws.row_dimensions[ws.max_row].outlineLevel = 1
    for col, w in zip("ABCDEFGHIJKL", (14, 10, 34, 12, 10, 12, 13, 14, 13, 10, 20, 44)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[4].number_format = "0.0000"
        row[5].number_format = "#,##0.0000"
        for i in (6, 7, 10):
            row[i].number_format = "#,##0.00;[Red]-#,##0.00"
        row[8].number_format = "0.00"


def writable_path(path):
    """Excel locks a workbook that is open, and Windows then refuses to overwrite it. If `path` is locked,
    return path_2.xlsx, path_3.xlsx, ... instead of failing."""
    base, ext = os.path.splitext(path)
    cand, n = path, 1
    while os.path.exists(cand):
        try:
            with open(cand, "ab"):
                break
        except PermissionError:
            n += 1
            cand = f"{base}_{n}{ext}"
    return cand


def autosize(ws, widths=None):
    for i, col in enumerate(ws.columns, 1):
        w = max((len(str(c.value)) if c.value is not None else 0) for c in col[:80])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 48)
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    ws.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser(description="US surge backtest -> Excel report grouped by ticker")
    ap.add_argument("--entry-mode", choices=["stop", "limit", "optimistic"], default="stop")
    ap.add_argument("--vol-mult", type=float, default=3.5)
    ap.add_argument("--capital", type=float, default=20000.0)
    ap.add_argument("--position", type=float, default=1000.0, help="base stake per slot")
    ap.add_argument("--slots", type=int, default=20)
    ap.add_argument("--from", dest="first", default="2024-11")
    ap.add_argument("--to", dest="last", default="2026-09")
    ap.add_argument("--data-through", default="2026-09-18", help="ignore later data (drops an unfinished day)")
    ap.add_argument("--include-hk", action="store_true", help="keep .HK tickers from the tracked list (HKD prices)")
    ap.add_argument("--out", help="output .xlsx path")
    args = ap.parse_args()

    tag = f"{args.entry_mode}_vm{args.vol_mult:g}_{int(args.slots)}slots"
    run_dir = os.path.join(OUT_ROOT, tag)
    out = args.out or os.path.join(OUT_ROOT, f"us_surge_backtest_{tag}.xlsx")

    conn = sqlite3.connect(study.DB_PATH)
    by_ticker = study.load_series_by_ticker(conn)
    conn.close()
    for t in list(by_ticker):
        if t.endswith(".HK") and not args.include_hk:       # e.g. HK funds in the tracked list: HKD prices, not US
            del by_ticker[t]
            continue
        by_ticker[t] = [r for r in by_ticker[t] if r["date"] <= args.data_through]
        study.add_ma21(by_ticker[t])
    print(f"US universe {len(by_ticker)} tickers | data through {args.data_through} | fills: {args.entry_mode} | volume x{args.vol_mult:g}")

    n_sig = generate_trades(by_ticker, args.entry_mode, args.vol_mult, args.first, args.last, os.path.join(run_dir, "strategies"))
    cwd = os.getcwd()
    os.chdir(run_dir)                                            # capital_simulation reads ./strategies/
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            res = capsim.run_simulation(args.capital, args.position, "strategies/capital_simulation_us.xlsx",
                                        compound_slots=args.slots)
    finally:
        os.chdir(cwd)
    s, trades, skipped, log, open_pos = res["summary"], res["trades"], res["skipped"], res["log"], res["open"]
    trades = trades.copy()
    trades["days_held"] = (pd.to_datetime(trades.final_sell_date) - pd.to_datetime(trades.buy_date)).dt.days
    trades["buy_month"] = trades.buy_date.astype(str).str[:7]

    # ---- ticker summary ----
    def agg(g):
        return pd.Series({"trades": len(g), "wins": int((g.total_pnl > 0).sum()), "losses": int((g.total_pnl <= 0).sum()),
                          "win_rate_%": round(100 * (g.total_pnl > 0).mean(), 1), "total_P/L_$": round(g.total_pnl.sum(), 2),
                          "avg_trade_%": round(g.total_pnl_pct.mean(), 2), "best_%": round(g.total_pnl_pct.max(), 2),
                          "worst_%": round(g.total_pnl_pct.min(), 2), "avg_days_held": round(g.days_held.mean(), 1),
                          "invested_$": round(g.invested.sum(), 2)})
    tsum = trades.groupby("ticker").apply(agg).reset_index().sort_values("total_P/L_$", ascending=False)

    # ---- summary ----
    monthly = trades.groupby("buy_month").agg(trades=("ticker", "count"), wins=("total_pnl", lambda x: int((x > 0).sum())),
                                              pnl=("total_pnl", lambda x: round(x.sum(), 2))).reset_index()
    monthly["win_rate_%"] = (100 * monthly.wins / monthly.trades).round(1)
    monthly = monthly.rename(columns={"buy_month": "buy month", "pnl": "P/L ($)"})
    r = trades.total_pnl_pct
    info = [
        ("STRATEGY", ""),
        ("Universe", f"US: tracked tickers + S&P 500 / Russell 1000 ({len(by_ticker)} tickers)"),
        ("Surge day", f"volume > {args.vol_mult:g}x its 10-day average, close > MA200, MA50 > MA200, up day"),
        ("Drop day", "first red candle (close < open) after the surge, within 10 trading days"),
        ("Buy", f"buy level = surge-day close, within 5 trading days after the drop day; fills: {args.entry_mode} "
                f"({'pay max(open, level) once the High reaches the level' if args.entry_mode == 'stop' else 'needs the Low to reach the level, pay min(open, level)' if args.entry_mode == 'limit' else 'OLD optimistic rule: fill at the level even if never traded'})"),
        ("Sell", "half at +15% (close vs surge close); the rest when a close is below MA21 and the next open is still below"),
        ("Period", f"surge days {args.first} .. {args.last}, price data through {args.data_through}"),
        ("Costs", "none modelled (US)"),
        ("", ""),
        ("ACCOUNT", ""),
        ("Start capital", f"${args.capital:,.0f}   ({args.slots} slots, base stake ${args.position:,.0f} per slot; each buy = base + an even share of the profit so far; never more than {args.slots} open positions)"),
        ("Signals found", f"{n_sig:,}"),
        ("Trades taken", f"{int(s['trades_executed'])}   (skipped for no free slot / cash: {int(s['trades_skipped'])})"),
        ("Ending value", f"${s['total_ending_value']:,.2f}   (cash ${s['ending_cash']:,.2f} + open positions ${s['value_in_open_positions']:,.2f})"),
        ("Total P/L", f"${s['total_pnl']:,.2f}   ({s['total_return_pct']:.1f}% on start capital)"),
        ("Win rate", f"{100 * (trades.total_pnl > 0).mean():.1f}%   ({int((trades.total_pnl > 0).sum())} wins / {int((trades.total_pnl <= 0).sum())} losses)"),
        ("Avg / median trade", f"{r.mean():.2f}% / {r.median():.2f}%"),
        ("Best / worst trade", f"{r.max():.1f}% / {r.min():.1f}%"),
        ("Still open", f"{int((trades.status != 'closed').sum())} position(s), marked at the last close"),
        ("", ""),
        ("HOW TO READ", ""),
        ("By Ticker", "one block per ticker: a bold header with the ticker total, then each trade's BUY / SELL rows. Click the +/- (or the outline 1/2 buttons top-left) to collapse or expand."),
        ("By Buy Date", "the same buy / sell detail with one block per trade, in the order the trades were bought; the header shows the running total of trade P/L."),
        ("Amount ($)", "negative = cash paid for a buy, positive = cash received from a sale"),
        ("Trade P/L", "shown on the last row of each trade: proceeds from all sells minus the amount paid"),
        ("Caveat", "daily bars only (intraday order of highs/lows unknown); no slippage or partial fills; backtest, not a forecast"),
    ]

    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    wanted = out
    out = writable_path(out)
    if out != wanted:
        print(f"note: {os.path.basename(wanted)} is open in Excel, so this run was saved as {os.path.basename(out)}")
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        wb = w.book
        ws = wb.create_sheet("Summary")
        for k, v in info:
            ws.append([k, v])
        for row in ws.iter_rows():
            if row[0].value in ("STRATEGY", "ACCOUNT", "HOW TO READ"):
                for c in row:
                    c.font, c.fill = HEAD_FONT, HEAD_FILL
            else:
                row[0].font = BOLD
        ws.append([])
        ws.append(["Monthly results (by buy month)"])
        ws.cell(ws.max_row, 1).font = BOLD
        start = ws.max_row + 1
        ws.append(list(monthly.columns))
        for c in ws[ws.max_row]:
            c.font, c.fill = HEAD_FONT, HEAD_FILL
        for row in monthly.itertuples(index=False):
            ws.append(list(row))
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 140
        for col in "CDEF":
            ws.column_dimensions[col].width = 12
        write_by_ticker(wb.create_sheet("By Ticker"), trades)
        write_by_buy_date(wb.create_sheet("By Buy Date"), trades)
        tsum.to_excel(w, sheet_name="Ticker Summary", index=False)
        cols = ["ticker", "surge_date", "drop_date", "buy_date", "buy_price", "shares_bought", "invested",
                "partial_sell_date", "partial_sell_price", "partial_sell_proceeds", "final_sell_date", "final_sell_price",
                "final_sell_proceeds", "exit_type", "status", "days_held", "total_proceeds", "total_pnl", "total_pnl_pct"]
        trades[cols].sort_values(["buy_date", "ticker"]).to_excel(w, sheet_name="Trades", index=False)
        log.to_excel(w, sheet_name="Transactions (chrono)", index=False)
        skipped.to_excel(w, sheet_name="Skipped Signals", index=False)
        open_pos.to_excel(w, sheet_name="Still Open", index=False)
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
        wb.move_sheet("Summary", offset=-wb.index(wb["Summary"]))
        for name in ("Ticker Summary", "Trades", "Transactions (chrono)", "Skipped Signals", "Still Open"):
            autosize(wb[name])

    print(f"signals {n_sig} | taken {int(s['trades_executed'])} | skipped {int(s['trades_skipped'])} | "
          f"P/L ${s['total_pnl']:,.0f} ({s['total_return_pct']:.1f}%) | win rate {100 * (trades.total_pnl > 0).mean():.1f}%")
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
